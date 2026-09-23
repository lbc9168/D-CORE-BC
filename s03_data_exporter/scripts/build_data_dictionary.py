"""build_data_dictionary.py — 生成 data_dictionary.xlsx（3 sheet：READ ME / Columns / Enums）。

单一可信源：Enums sheet 从根 enums_mapping.json 生成；README 只描述稳定的数据粒度、
字段口径和使用约定，不嵌入随数据增量变化的统计。列说明为文档散文，维护在本脚本
_COLUMNS 表中——events_exporter._COLUMNS 加列时必须同步这里
（脚本会校验两边列集合一致，漏列即报错退出）。

运行（仓库根）：uv run --no-sync python s03_data_exporter/scripts/build_data_dictionary.py
输出：s04_web_visualization/public/data/data_dictionary.xlsx（git 跟踪，随代码提交）
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_ROOT = next(p for p in Path(__file__).resolve().parents if (p / "bri_enums.py").exists())
_DATA_DIR = _ROOT / "s04_web_visualization" / "public" / "data"
_OUT = _DATA_DIR / "data_dictionary.xlsx"
sys.path.insert(0, str(_ROOT / "s03_data_exporter" / "src"))
from s03_data_exporter import events_exporter  # noqa: E402

# ── 样式常量 ──────────────────────────────────────────────────────────────────
_BRAND = "2F6562"
_WARN_FILL = PatternFill("solid", fgColor="FFF7E6")
_HEAD_FILL = PatternFill("solid", fgColor="E0EEEE")
_WRAP = Alignment(wrap_text=True, vertical="top")


# ── Sheet 1: READ ME ─────────────────────────────────────────────────────────

def _build_readme(wb: Workbook) -> None:
    """READ ME sheet：给人和 AI 的使用说明——粒度陷阱/容量语义/列适用矩阵/通用约定。"""
    ws = wb.active
    ws.title = "READ ME"
    ws.column_dimensions["A"].width = 118

    def line(text: str = "", *, bold: bool = False, size: int = 10, warn: bool = False, color: str | None = None):
        ws.append([text])
        c = ws.cell(row=ws.max_row, column=1)
        c.font = Font(bold=bold, size=size, color=color)
        c.alignment = _WRAP
        if warn:
            c.fill = _WARN_FILL

    line("D-CORE events_export.csv — READ THIS FIRST (for humans & AI analysts)", bold=True, size=14, color=_BRAND)
    line("Dataset scope: Chinese overseas renewable-energy and related projects. This workbook defines the CSV "
         "grain, fields, applicability, units and controlled values.")
    line("This sheet lists the pitfalls that WILL produce wrong numbers if ignored. "
         "Read the four numbered signals below before any aggregation.", bold=True, warn=True)
    line()

    line("1. GRAIN & DE-DUPLICATION — one row is NOT one project", bold=True, size=11, color=_BRAND)
    line("• Grain: one row = one milestone event × one Chinese entity. A project (canonical_id) has 1..N events "
         "(event_id); an event involving N Chinese entities is repeated N times (entity columns differ, all else repeats).")
    line("• Project-level values (capacity_mw, total_investment_usd, storage_*, pv_capacity_mw, amounts) REPEAT on "
         "every row of the same project. Summing them over all rows double-counts.", warn=True)
    line("• Recipe — count projects: df['canonical_id'].nunique().")
    line("• Recipe — sum capacity: df.sort_values('event_id').drop_duplicates('canonical_id', keep='last')"
         "['capacity_mw'].sum()  (dedupe to one row per project first; keep='last' takes the latest event).")
    line("• Use event_id + year for time-series (events are dated); use canonical_id for anything project-level.")
    line()

    line("2. CAPACITY & ENERGY-MIX SEMANTICS — keep generation, storage and manufacturing metrics separate", bold=True, size=11, color=_BRAND)
    line("• capacity_mw = installed POWER-GENERATION capacity in MW. It is blank by design for Energy Storage, "
         "Transmission and Manufacturing projects (their scale is measured differently, see signal 3).", warn=True)
    line("• Energy storage uses TWO fields: storage_power_mw (MW, how fast it charges/discharges) and "
         "storage_energy_mwh (MWh, how much it stores). NEVER add MW to MWh, and never add storage MW into "
         "generation capacity_mw sums.", warn=True)
    line("• Integrated Energy means a mixed generation system (renewable and/or thermal sources). capacity_mw is "
         "total generation capacity; pv_capacity_mw is a separately captured PV component when available; "
         "storage_* are separate storage metrics. Do not infer missing component capacities from names or totals.")
    line("• generation_sources stores pipe-separated atomic generation-source keys; storage is never a generation "
         "source. storage_status records evidence strength, while technology_tags_cn/_en provides controlled "
         "combination labels such as Solar + Storage, Wind + Storage or Wind + Solar + Storage.")
    line("• storage_* fields may appear on Solar PV, Wind Power, Integrated Energy or Energy Storage projects. "
         "Filter by field presence or storage_status/technology tags, not only by category.")
    line()

    line("3. COLUMN APPLICABILITY BY CATEGORY — which columns mean anything for which projects", bold=True, size=11, color=_BRAND)
    matrix = [
        ("Column group", "Generation (PV/Wind/Hydro/Geo/Biomass)", "Energy Storage", "Integrated Energy (hybrid)", "Transmission", "Manufacturing"),
        ("capacity_mw", "✓", "", "✓ (plant total)", "", ""),
        ("subcategory_cn / _en", "✓ Solar PV only (Utility-scale/Distributed)", "", "", "", ""),
        ("generation_sources", "✓ atomic generation source(s)", "", "✓ mixed generation source(s)", "", ""),
        ("storage_status / technology_tags_*", "if storage/combination disclosed", "✓", "if storage/combination disclosed", "", ""),
        ("pv_capacity_mw", "if separately captured", "", "if separately captured", "", ""),
        ("storage_power_mw / storage_energy_mwh", "if disclosed", "✓", "✓ (storage sub-component)", "", ""),
        ("voltage_kv / line_length_km / transmission_capacity_mw", "", "", "", "✓", ""),
        ("production_capacity_value / _unit", "", "", "", "", "✓"),
        ("stage / year / month / entities / amounts", "✓", "✓", "✓", "✓", "✓"),
    ]
    start = ws.max_row + 1
    for r in matrix:
        ws.append(list(r))
    for col in range(1, 7):
        cell = ws.cell(row=start, column=col)
        cell.font = Font(bold=True, size=9)
        cell.fill = _HEAD_FILL
        cell.alignment = _WRAP
    for row in ws.iter_rows(min_row=start + 1, max_row=ws.max_row, max_col=6):
        for cell in row:
            cell.font = Font(size=9)
            cell.alignment = _WRAP
    for col, w in zip("BCDEF", (26, 14, 22, 13, 13)):
        ws.column_dimensions[col].width = w
    line("• Blank cells above = the column is not applicable to that category (structurally empty, not missing data).")
    line("• Manufacturing capacities use HETEROGENEOUS units (GW/year, GWh/year, t/year). "
         "Values originally in 10,000 tonnes/year are multiplied by 10,000 for t/year. "
         "Never sum production_capacity_value across different units.", warn=True)
    line()

    line("4. GENERAL CONVENTIONS", bold=True, size=11, color=_BRAND)
    line("• Blank = not disclosed in sources, NOT zero.")
    line("• All monetary fields are in MILLION USD. contract_amount_usd / financing_amount_usd are only collected "
         "for events from July 2026 onward; earlier events are blank by policy (not missing).", warn=True)
    line("• is_central_soe describes the Chinese ENTITY on that row (1/0), not the project; a project may mix "
         "SOE and non-SOE rows.")
    line("• Scope baseline: milestone events from 2022 onward; projects whose latest event predates 2022 are "
         "excluded; projects located in mainland China / Hong Kong / Macau / Taiwan are excluded.")
    line("• Row order is for human browsing only (generation+storage block, then transmission, then manufacturing; "
         "sorted by canonical_id, event_id inside each block). Do not rely on row order programmatically.")
    line("• Stage vocabulary: fine-grained stage_cn/_en roll up to four macro groups (stage_group_*): "
         "Pre-construction / In Construction / Construction Completed / Acquired — see the Enums sheet.")


# ── Sheet 2: Columns ─────────────────────────────────────────────────────────

# (column, description, applies_to, example) —— 与 events_exporter._COLUMNS 同步（脚本校验）
_COLUMNS = [
    ("canonical_id", "Physical project ID - one per real-world project. Format {ISO3}-{CATEGORY}-{NNNN}.", "All", "VNM-WD-0001"),
    ("event_id", "Milestone (event) ID - one per milestone of a project. Format {canonical_id}-{NN}.", "All", "VNM-WD-0001-01"),
    ("project_name_en", "Project name in English.", "All", "Vietnam 128MW Nearshore Wind Farm Project"),
    ("project_name_cn", "Project name in Chinese.", "All", "越南128MW近海风电项目"),
    ("country_en", "Host country name in English.", "All", "Vietnam"),
    ("country_iso", "Host country ISO 3166-1 alpha-3 code.", "All", "VNM"),
    ("province", "First-level administrative region (province/state), if known; blank otherwise.", "All", "Binh Thuan"),
    ("category_en", "Energy category in English (Solar PV / Wind Power / Hydropower / Geothermal & Waste Heat / Biomass / Energy Storage / Integrated Energy / Transmission & Distribution / New Energy Manufacturing).", "All", "Wind Power"),
    ("category_cn", "Energy category in Chinese.", "All", "风电"),
    ("subcategory_cn", "Project-form subcategory in Chinese. Solar PV only: 集中式 (utility-scale; default when unspecified) / 分布式 (distributed). Energy combinations are stored in technology_tags_cn/_en, not subcategory.", "Solar PV only", "集中式"),
    ("subcategory_en", "Project-form subcategory in English (Utility-scale / Distributed). Energy combinations are stored in technology_tags_cn/_en.", "Solar PV only", "Utility-scale"),
    ("generation_sources", "Event-level pipe-separated atomic generation-source keys in Chinese (e.g. 光伏 | 风电). Energy storage is excluded by definition; blank means no structured source fact was determined.", "Generation & Integrated Energy", "光伏 | 风电"),
    ("storage_status", "Controlled storage-evidence status: confirmed / planned / optional / unclear / not_mentioned. Only confirmed storage exports numeric storage_* values.", "Generation, Integrated Energy & Energy Storage", "confirmed"),
    ("technology_tags_cn", "Pipe-separated controlled Chinese technology-combination labels (光储 / 风储 / 风光 / 风光储). These project-level labels repeat for every event and entity and are independent of project-form subcategory.", "Generation & Integrated Energy", "风光储"),
    ("technology_tags_en", "English rendering of technology_tags_cn (Solar + Storage / Wind + Storage / Wind + Solar / Wind + Solar + Storage).", "Generation & Integrated Energy", "Wind + Solar + Storage"),
    ("stage_cn", "Milestone stage in Chinese (签约 / 中标 / 开工 / 在建 / 竣工 / 投入运营 / 收购 ...).", "All", "签约"),
    ("stage_en", "Milestone stage in English (Signed / Awarded / Construction Started / Under Construction / Completed / Operational / Acquired ...).", "All", "Signed"),
    ("stage_group_cn", "Milestone macro-group in Chinese (前期 / 在建 / 完工 / 收购).", "All", "前期"),
    ("stage_group_en", "Milestone macro-group in English (Pre-construction / In Construction / Construction Completed / Acquired).", "All", "Pre-construction"),
    ("year", "Year of this milestone event.", "All", "2026"),
    ("month", "Month of this milestone event (YYYY-MM), if known; blank otherwise.", "All", "2026-02"),
    ("capacity_mw", "Installed power-generation capacity in MW. Blank by design for Energy Storage / Transmission / Manufacturing (not measured in generation MW).", "Generation & Integrated Energy", "128.0"),
    ("pv_capacity_mw", "Separately captured Solar PV component capacity in MW, when disclosed. Most common in solar+storage and Integrated Energy projects; blank otherwise.", "Any project with a separately captured PV component", "100.0"),
    ("storage_power_mw", "Energy storage power rating in MW, if disclosed; blank otherwise. Not limited to the Energy Storage category.", "Any project disclosing storage", "50.0"),
    ("storage_energy_mwh", "Energy storage capacity (energy) in MWh, if disclosed; blank otherwise. Not limited to the Energy Storage category.", "Any project disclosing storage", "200.0"),
    ("voltage_kv", "Transmission line/substation voltage level in kV. A grade, not a quantity - do not sum.", "Transmission only", "400.0"),
    ("line_length_km", "Transmission line length in km, if disclosed.", "Transmission only", "270.0"),
    ("transmission_capacity_mw", "Transmission capacity in MW, if disclosed (sparse).", "Transmission only", "5000.0"),
    ("production_capacity_value", "Manufacturing production capacity value; unit given in production_capacity_unit. Units are heterogeneous - never sum across units.", "Manufacturing only", "5.0"),
    ("production_capacity_unit", "Unit of production_capacity_value (GW/year; GWh/year; t/year = tonnes per year; Unknown). Values originally in 10,000 tonnes/year are multiplied by 10,000.", "Manufacturing only", "GW/year"),
    ("total_investment_usd", "Total project investment in million USD, if disclosed; blank otherwise.", "All", "250"),
    ("contract_amount_usd", "Contract value (EPC / general contracting / construction / equipment supply) in million USD. Only collected for events from 2026-07 onward; earlier events blank by policy.", "All (2026-07+ events)", "167"),
    ("financing_amount_usd", "Debt financing amount (loan / credit facility / syndicated loan / export credit) in million USD. Only collected for events from 2026-07 onward; earlier events blank by policy.", "All (2026-07+ events)", "486"),
    ("financing_model_en", "Financing / delivery model in English (e.g. EPC, BOT, PPP); 'Unknown' if not disclosed.", "All", "EPC"),
    ("developer_cn", "Project developer / owner in Chinese, if disclosed for this event.", "All", "越南REE集团"),
    ("developer_en", "Project developer / owner in English; blank if not yet translated.", "All", "Vietnam REE Corporation"),
    ("entity_cn", "Chinese entity involved in this event, in Chinese. One row per entity (see READ ME signal 1).", "All", "中国能建国际集团"),
    ("entity_en", "Chinese entity involved in this event, in English.", "All", "China Energy Engineering Group Co., Ltd. (CEEC)"),
    ("subsidiary_cn", "Subsidiary of the Chinese entity in Chinese, if any; blank otherwise.", "All", "中国电力工程顾问集团有限公司"),
    ("subsidiary_en", "Subsidiary of the Chinese entity in English, if any; blank otherwise.", "All", "China Power Engineering Consulting Group Co., Ltd. (CPECC)"),
    ("involvement_type_cn", "How the Chinese entity is involved, in Chinese (业主 / 总承包 / 设备供应商 / 投资方 ...).", "All", "设计/施工服务方"),
    ("involvement_type_en", "How the Chinese entity is involved, in English (Owner / EPC Contractor / Equipment Supplier ...).", "All", "Design/Construction Service"),
    ("is_central_soe", "Whether THIS Chinese entity is a central state-owned enterprise. 1 = yes, 0 = no; blank when no entity on the row.", "All", "1"),
]


def _build_columns(wb: Workbook) -> None:
    """Columns sheet：全列字典（Column / Description / Applies to / Example）。"""
    ws = wb.create_sheet("Columns")
    ws.append(["Column", "Description", "Applies to", "Example"])
    for col in range(1, 5):
        c = ws.cell(row=1, column=col)
        c.font = Font(bold=True)
        c.fill = _HEAD_FILL
    for row in _COLUMNS:
        ws.append(list(row))
    for row in ws.iter_rows(min_row=2, max_col=4):
        for cell in row:
            cell.alignment = _WRAP
            cell.font = Font(size=9)
    for col, w in zip("ABCD", (26, 80, 26, 34)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


# ── Sheet 3: Enums ───────────────────────────────────────────────────────────

def _build_enums(wb: Workbook) -> None:
    """Enums sheet：类别/阶段映射/子类/制造业单位——全部从根 enums_mapping.json 生成。"""
    enums = json.loads((_ROOT / "enums_mapping.json").read_text(encoding="utf-8"))
    ws = wb.create_sheet("Enums")
    ws.column_dimensions["A"].width = 24
    for col, w in zip("BCDEF", (30, 16, 40, 20, 16)):
        ws.column_dimensions[col].width = w

    def header(title: str):
        ws.append([])
        ws.append([title])
        c = ws.cell(row=ws.max_row, column=1)
        c.font = Font(bold=True, size=11, color=_BRAND)

    def table(rows: list[tuple], head: tuple):
        ws.append(list(head))
        for col in range(1, len(head) + 1):
            c = ws.cell(row=ws.max_row, column=col)
            c.font = Font(bold=True, size=9)
            c.fill = _HEAD_FILL
        for r in rows:
            ws.append(list(r))
            for col in range(1, len(r) + 1):
                cell = ws.cell(row=ws.max_row, column=col)
                cell.font = Font(size=9)
                cell.alignment = _WRAP

    # 类别表（含是否进 CSV：hidden 且非 TX/MF 的核电/火电不导出）
    basis = {
        "ES": "storage_power_mw (MW) + storage_energy_mwh (MWh)",
        "HB": "total generation capacity_mw (MW), plus disclosed pv/storage sub-metrics",
        "TX": "voltage_kv / line_length_km / transmission_capacity_mw",
        "MF": "production_capacity_value + unit (heterogeneous)",
    }
    cat_rows = []
    for zh, m in enums["category"].items():
        if zh == "Unknown":
            continue
        code = m["code"]
        in_csv = "No (excluded)" if (m.get("hidden") and code not in ("TX", "MF")) else "Yes"
        cat_rows.append((code, m["en"], zh, basis.get(code, "capacity_mw (MW)"), in_csv))
    header("Categories (scale-metric basis & CSV inclusion)")
    table(cat_rows, ("Code", "English", "Chinese", "Scale metric", "In CSV?"))

    # 阶段 → 4 大类映射
    vg = enums["viz_group"]
    stage_rows = []
    for zh, m in enums["project_stage"].items():
        if zh == "Unknown":
            continue
        g = m.get("viz_group")
        gm = vg.get(g) if g else None
        stage_rows.append((zh, m["en"], gm["zh"] if gm else "—", gm["en"] if gm else "— (not shown)"))
    header("Milestone stages → macro groups (stage_group_*)")
    table(stage_rows, ("Stage (zh)", "Stage (en)", "Group (zh)", "Group (en)"))

    # 子类
    sub_rows = []
    for cat, d in enums["subcategory"].items():
        if cat.startswith("_"):
            continue
        for v, m in d["values"].items():
            note = "default when unspecified" if v == d.get("_default") else ""
            sub_rows.append((cat, v, m["en"], note))
    header("Project-form subcategories (Solar PV only)")
    table(sub_rows, ("Category (zh)", "Value (zh)", "English", "Note"))

    # 发电源原子事实（储能不属于发电源）
    source_rows = [
        (zh, m["en"], m.get("order", ""))
        for zh, m in sorted(enums["generation_source"].items(), key=lambda item: item[1].get("order", 999))
    ]
    header("Atomic generation sources (energy storage is excluded)")
    table(source_rows, ("Key (zh)", "English", "Order"))

    # 储能证据状态
    storage_rows = [
        (key, m["zh"], m["en"], m.get("priority", ""))
        for key, m in sorted(enums["storage_status"].items(), key=lambda item: item[1].get("priority", 0), reverse=True)
    ]
    header("Storage evidence status")
    table(storage_rows, ("Key", "Chinese", "English", "Priority"))

    # 技术组合标签（独立于项目形态子类）
    tag_rows = [
        (key, m["en"], " | ".join(m.get("allowed_categories", [])))
        for key, m in enums["technology_tag"].items()
    ]
    header("Technology combination tags (independent of subcategory)")
    table(tag_rows, ("Key (zh)", "English", "Allowed project categories"))

    # 制造业产能单位
    header("Manufacturing production-capacity units (never sum across units)")
    unit_rows = [
        ("GW/year", "GW per year (module/cell manufacturing lines)"),
        ("GWh/year", "GWh per year (battery factories)"),
        ("t/year", "Tonnes per year (polysilicon/materials); source 10k-tonne values multiplied by 10,000"),
        ("Unknown", "value disclosed without a usable unit"),
    ]
    table(unit_rows, ("Unit", "Meaning"))


def main() -> None:
    """生成 3-sheet 数据字典；先校验列集合与 events_exporter 同步。"""
    doc_cols = [c[0] for c in _COLUMNS]
    if doc_cols != list(events_exporter._COLUMNS):
        missing = set(events_exporter._COLUMNS) ^ set(doc_cols)
        raise SystemExit(f"字典列与 events_exporter._COLUMNS 不同步（差集 {missing}），请先更新本脚本")
    wb = Workbook()
    _build_readme(wb)
    _build_columns(wb)
    _build_enums(wb)
    _OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(_OUT)
    print(f"✓ data_dictionary.xlsx 已生成 → {_OUT}")


if __name__ == "__main__":
    main()
